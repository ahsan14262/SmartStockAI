from io import BytesIO
import pandas as pd
from app.core.security import neutralize_spreadsheet_formula
def csv_bytes(df):
    x=df.map(neutralize_spreadsheet_formula) if hasattr(df,"map") else df
    return x.to_csv(index=False).encode("utf-8")
def xlsx_bytes(df,sheet="Report"):
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as w: df.to_excel(w,index=False,sheet_name=sheet[:31])
    return out.getvalue()
def docx_bytes(title,paragraphs):
    from docx import Document
    d=Document(); d.add_heading(title,0)
    for p in paragraphs:d.add_paragraph(str(p))
    out=BytesIO(); d.save(out); return out.getvalue()
def pdf_bytes(title,lines):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    out=BytesIO(); styles=getSampleStyleSheet(); story=[Paragraph(title,styles["Title"])]
    for x in lines: story += [Paragraph(str(x),styles["BodyText"]),Spacer(1,6)]
    SimpleDocTemplate(out).build(story); return out.getvalue()
